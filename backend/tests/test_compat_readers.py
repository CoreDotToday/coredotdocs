"""독립 리더 호환성: 우리가 쓴 DOCX를 python-docx(제3자 라이브러리)가 읽는지.

HWPX는 rhwp(`tests/test_rhwp_node.py`)와 rhwp-studio로 확인한다. 여기서는 DOCX를 자동화한다.
"""

from __future__ import annotations

import docx as python_docx

from app.engine import commands
from app.engine.docx import create_blank
from tests.test_commands_full import png

MD = "# 호환성 검증 문서\n\n이 문단은 **굵게**와 *기울임*을 포함한 본문입니다.\n\n- 목록 하나\n\n| 구분 | 내용 |\n|---|---|\n| 가 | 나 |\n\n<!-- pagebreak -->\n둘째 쪽"


def test_python_docx_reads_everything_we_write(tmp_path):
    doc = create_blank("호환성 검증")
    commands.insert_content(doc, MD)
    commands.insert_image(doc, png(64, 32), width_mm=30)
    commands.format_text(doc, "본문", color="#C00000", size=14, underline=True)
    commands.format_paragraphs(doc, 2, align="center", line_spacing=150)
    commands.insert_table_rows(doc, 1, rows=[["새", "행"]])
    commands.set_table_cells(doc, 1, [{"row": 1, "col": 1, "format": {"bold": True}}])
    path = tmp_path / "compat.docx"
    path.write_bytes(doc.to_bytes())

    d = python_docx.Document(str(path))
    texts = [p.text for p in d.paragraphs]
    assert texts[1] == "호환성 검증 문서" and texts[2] == "이 문단은 굵게와 기울임을 포함한 본문입니다." and "둘째 쪽" in texts
    assert d.paragraphs[1].style.name == "Heading 1"
    assert [[c.text for c in r.cells] for r in d.tables[0].rows] == [["구분", "내용"], ["가", "나"], ["새", "행"]]
    assert len(d.inline_shapes) == 1 and round(d.inline_shapes[0].width.mm) == 30
    run = next(r for p in d.paragraphs for r in p.runs if "본문" in r.text)
    assert str(run.font.color.rgb) == "C00000" and run.font.size.pt == 14 and run.font.underline is True
    assert d.paragraphs[1].paragraph_format.alignment == 1  # center
    assert round(d.sections[0].page_width.mm) == 210
    header_run = d.tables[0].rows[0].cells[0].paragraphs[0].runs[0]
    assert header_run.font.bold is True
